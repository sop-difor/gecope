#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_stg_orse.py
=================
Extrator do ORSE para o modelo NOVO (pos-reestruturacao de 2026-09).

Le os arquivos de custo de referencia do ORSE (.xls / .xlsx) de UM mes -- o de
SERVICOS e o de INSUMOS -- e gera **um unico** CSV no formato da area de
recebimento `stg_orse`:

    identificacao,codigo,descricao,unidade,preco_unitario,tipo_encargo,referencia

- `identificacao`: 'C' para o arquivo de SERVICOS, 'I' para o de INSUMOS
  (detectado pelo nome do arquivo; da para forcar com --tipo).
- `tipo_encargo`: sempre `onerada` (ORSE nao tem desonerada).
- `referencia`: `AAAA-MM-01` -- vem de --mes, ou do nome da pasta/arquivo.
- Codigo de insumo que o Excel entrega como float (`8792.0`) vira `8792`
  (a convencao historica no banco e inteiro sem `.0`).

Depois: subir o CSV com `subir_stg_orse.mjs` (ou pela tela do Supabase) e o
banco roda `rt_aplicar_orse('AAAA-MM-01')`, que grava so o delta.

NAO escreve no banco. NAO gera INSERTs. (O antigo `motor_orse.py`, que gravava
direto em `public.orse_itens`, ficou obsoleto: `orse_itens` hoje e uma VIEW.)

Uso
---
# um mes, arquivos explicitos:
python gerar_stg_orse.py --mes 2026-04 \
    --servicos ".../ORSE_Custo_ref_SERVICOS_ABRIL_2026.xls" \
    --insumos  ".../ORSE_Custo_ref_INSUMOS_ABRIL_2026.xls"

# um mes, apontando a pasta (acha SERVICOS/INSUMOS e o mes pelos nomes):
python gerar_stg_orse.py --pasta ".../ORSE_2026/04 - ABRIL"

# so servicos (quando ainda nao ha planilha de insumos):
python gerar_stg_orse.py --mes 2026-06 --servicos ".../....xls"

Parsing adaptado de motor_orse.py (mesmas regras de cabecalho/preco/referencia).
"""

import argparse
import csv
import glob
import os
import re
import sys
import unicodedata
from datetime import date

MESES_PT = {
    "JANEIRO": 1, "FEVEREIRO": 2, "MARCO": 3, "ABRIL": 4,
    "MAIO": 5, "JUNHO": 6, "JULHO": 7, "AGOSTO": 8,
    "SETEMBRO": 9, "OUTUBRO": 10, "NOVEMBRO": 11, "DEZEMBRO": 12,
}

COLUNAS = ["identificacao", "codigo", "descricao", "unidade",
           "preco_unitario", "tipo_encargo", "referencia"]


def _sem_acento(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in nfkd if not unicodedata.combining(c)).upper().strip()


def parse_preco(valor):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return round(float(valor), 2)
    s = str(valor).strip()
    if s == "":
        return None
    s = s.replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def parse_referencia(texto_data: str):
    if not texto_data:
        return None
    t = _sem_acento(texto_data)
    m = re.search(r"(?<!\d)(\d{1,2})[/\-](\d{4})(?!\d)", t)
    if m:
        mes, ano = int(m.group(1)), int(m.group(2))
        if 1 <= mes <= 12:
            return date(ano, mes, 1)
    m = re.search(r"(?<!\d)(20\d{2})[_/\-](\d{1,2})(?!\d)", t)  # AAAA_MM (nome de arquivo)
    if m:
        ano, mes = int(m.group(1)), int(m.group(2))
        if 1 <= mes <= 12:
            return date(ano, mes, 1)
    ano_m = re.search(r"(?<!\d)(20\d{2})(?!\d)", t)
    if ano_m:
        ano = int(ano_m.group(1))
        for nome, num in MESES_PT.items():
            if nome in t:
                return date(ano, num, 1)
    return None


def normaliza_codigo(bruto) -> str:
    """'8792.0' -> '8792'; 'S03126' -> 'S03126'; espacos nas pontas fora."""
    s = str(bruto).strip()
    if re.fullmatch(r"\d+\.0+", s):
        return s.split(".")[0]
    return s


def ler_linhas(caminho: str):
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            sys.exit("Erro: instale 'xlrd' para ler .xls  ->  python -m pip install xlrd")
        book = xlrd.open_workbook(caminho)
        sh = book.sheet_by_index(0)
        return [sh.row_values(r) for r in range(sh.nrows)]
    elif ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("Erro: instale 'openpyxl' para ler .xlsx  ->  python -m pip install openpyxl")
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[("" if c is None else c) for c in row]
                for row in ws.iter_rows(values_only=True)]
    else:
        sys.exit(f"Extensao nao suportada: {ext}. Use .xls ou .xlsx")


def localizar_cabecalho(linhas):
    for i, linha in enumerate(linhas):
        norm = [_sem_acento(c) for c in linha]
        tem_codigo = any(x.startswith("CODIGO") for x in norm)
        tem_desc = any(x.startswith("DESCRICAO") for x in norm)
        if tem_codigo and tem_desc:
            mapa = {}
            for j, x in enumerate(norm):
                if x.startswith("CODIGO") and "codigo" not in mapa:
                    mapa["codigo"] = j
                elif x.startswith("DESCRICAO") and "descricao" not in mapa:
                    mapa["descricao"] = j
                elif x.startswith("UNIDADE") and "unidade" not in mapa:
                    mapa["unidade"] = j
                elif (x.startswith("PRECO") or x.startswith("CUSTO")) and "preco" not in mapa:
                    mapa["preco"] = j
                elif x.startswith("DATA") and "data" not in mapa:
                    mapa["data"] = j
            return i, mapa
    sys.exit("Erro: nao encontrei a linha de cabecalho (CODIGO/DESCRICAO) no arquivo.")


def extrair(caminho, identificacao, ref_fixa):
    linhas = ler_linhas(caminho)
    idx_cab, cols = localizar_cabecalho(linhas)
    c_cod, c_desc = cols["codigo"], cols["descricao"]
    c_unid, c_preco, c_data = cols.get("unidade"), cols.get("preco"), cols.get("data")

    ref_planilha = ref_fixa
    if ref_planilha is None:
        for i in range(0, idx_cab):
            r = parse_referencia(" ".join(str(c) for c in linhas[i]))
            if r:
                ref_planilha = r
                break
    if ref_planilha is None:
        ref_planilha = parse_referencia(os.path.basename(caminho))

    regs, sem_preco, sem_ref, vazias = [], 0, 0, 0
    for linha in linhas[idx_cab + 1:]:
        codigo = normaliza_codigo(linha[c_cod]) if c_cod < len(linha) else ""
        descricao = str(linha[c_desc]).strip() if c_desc < len(linha) else ""
        if codigo == "":
            vazias += 1
            continue

        unidade = None
        if c_unid is not None and c_unid < len(linha):
            unidade = str(linha[c_unid]).strip() or None

        preco = None
        if c_preco is not None and c_preco < len(linha):
            preco = parse_preco(linha[c_preco])
        if preco is None:
            sem_preco += 1

        ref = None
        if c_data is not None and c_data < len(linha):
            ref = parse_referencia(str(linha[c_data]))
        ref = ref or ref_planilha
        if ref is None:
            sem_ref += 1
            continue

        regs.append({
            "identificacao": identificacao,
            "codigo": codigo,
            "descricao": descricao,
            "unidade": unidade,
            "preco_unitario": preco,
            "tipo_encargo": "onerada",
            "referencia": ref,
        })
    return regs, {"sem_preco": sem_preco, "sem_ref": sem_ref, "vazias": vazias}


def tipo_por_nome(caminho):
    n = _sem_acento(os.path.basename(caminho))
    if "INSUMO" in n:
        return "I"
    if "SERVICO" in n or "SERVI" in n:
        return "C"
    return None


def achar_na_pasta(pasta):
    arqs = []
    for pad in ("*.xls", "*.xlsx", "*.XLS", "*.XLSX"):
        arqs += glob.glob(os.path.join(pasta, pad))
    arqs = [a for a in arqs if not os.path.basename(a).startswith("~$")]
    serv = [a for a in arqs if tipo_por_nome(a) == "C"]
    insu = [a for a in arqs if tipo_por_nome(a) == "I"]
    return (serv[0] if serv else None), (insu[0] if insu else None)


def main():
    p = argparse.ArgumentParser(description="Gera stg_orse_AAAA-MM.csv (servicos + insumos de um mes).")
    p.add_argument("--mes", help="AAAA-MM (referencia do mes). Se omitido, sai do nome da pasta/arquivo.")
    p.add_argument("--servicos", help="Planilha de SERVICOS (.xls/.xlsx)")
    p.add_argument("--insumos", help="Planilha de INSUMOS (.xls/.xlsx)")
    p.add_argument("--pasta", help="Pasta com as planilhas do mes (acha SERVICOS/INSUMOS)")
    p.add_argument("--tipo", choices=["C", "I"], help="Forca a identificacao quando so passa 1 arquivo")
    p.add_argument("--saida", default=".", help="Pasta de saida do CSV (padrao: pasta atual)")
    args = p.parse_args()

    serv, insu = args.servicos, args.insumos
    if args.pasta:
        s2, i2 = achar_na_pasta(args.pasta)
        serv = serv or s2
        insu = insu or i2
    if not serv and not insu:
        sys.exit("Nada para ler. Use --servicos/--insumos ou --pasta.")

    ref = None
    if args.mes:
        try:
            y, m = map(int, args.mes.split("-")[:2])
            ref = date(y, m, 1)
        except Exception:
            sys.exit("--mes deve ser AAAA-MM")
    if ref is None:
        for cand in (args.pasta, serv, insu):
            if cand:
                ref = parse_referencia(os.path.basename(str(cand).rstrip("/\\")))
                if ref:
                    break
    if ref is None:
        sys.exit("Nao consegui determinar o mes. Passe --mes AAAA-MM.")

    entradas = []
    if serv:
        entradas.append((serv, args.tipo or tipo_por_nome(serv) or "C"))
    if insu:
        entradas.append((insu, args.tipo or tipo_por_nome(insu) or "I"))

    todos, por_ident = [], {}
    for caminho, ident in entradas:
        regs, rel = extrair(caminho, ident, ref)
        print(f"  {os.path.basename(caminho)}  [{ident}]  "
              f"linhas={len(regs)}  sem_preco={rel['sem_preco']}  "
              f"sem_ref={rel['sem_ref']}  ignoradas={rel['vazias']}")
        todos += regs
        por_ident[ident] = por_ident.get(ident, 0) + len(regs)

    # dedup: (identificacao, codigo). Exato -> silencioso; divergente -> avisa e mantem o 1o.
    vistos, saida, dups_div = {}, [], 0
    for r in todos:
        k = (r["identificacao"], r["codigo"])
        if k in vistos:
            if vistos[k] != (r["descricao"], r["unidade"], r["preco_unitario"]):
                dups_div += 1
            continue
        vistos[k] = (r["descricao"], r["unidade"], r["preco_unitario"])
        saida.append(r)
    if dups_div:
        print(f"  aviso: {dups_div} codigo(s) repetido(s) no mesmo arquivo com valores diferentes (mantido o 1o).")

    os.makedirs(args.saida, exist_ok=True)
    destino = os.path.join(args.saida, f"stg_orse_{ref.isoformat()[:7]}.csv")
    with open(destino, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUNAS)
        for r in saida:
            w.writerow([
                r["identificacao"], r["codigo"], r["descricao"],
                "" if r["unidade"] is None else r["unidade"],
                "" if r["preco_unitario"] is None else f'{r["preco_unitario"]:.2f}',
                r["tipo_encargo"], r["referencia"].isoformat(),
            ])

    print(f"\nOK -> {destino}")
    print(f"     {len(saida)} linhas  (C={por_ident.get('C', 0)}  I={por_ident.get('I', 0)})  ref={ref.isoformat()}")
    if "I" not in por_ident:
        print("     ATENCAO: sem planilha de INSUMOS neste mes.")


if __name__ == "__main__":
    main()
